from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import sys

root = Path(__file__).resolve().parents[1]
for source in sorted((root / 'docx').glob('*.md')):
    if len(sys.argv) > 1 and not source.name.startswith(sys.argv[1]):
        continue
    doc = Document()
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.9)
    sec.left_margin = sec.right_margin = Cm(2.1)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(25)
    doc.styles['Title'].font.color.rgb = RGBColor.from_string('111111')
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 1'].font.color.rgb = RGBColor.from_string('137B7D')
    first = True
    for line in source.read_text(encoding='utf-8').splitlines():
        if not line.strip(): continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            if not first: doc.add_page_break()
            first = False
            doc.add_paragraph(line[3:], 'Heading 1')
        else:
            doc.add_paragraph(line)
    doc.core_properties.title = source.stem
    doc.core_properties.author = 'Codefront'
    dest = source.with_suffix('.docx')
    doc.save(dest)
    print(dest)
